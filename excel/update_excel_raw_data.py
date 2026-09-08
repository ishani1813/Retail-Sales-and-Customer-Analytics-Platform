import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

orders = pd.read_csv("data/processed/orders.csv", parse_dates=["order_date", "ship_date"])
customers = pd.read_csv("data/processed/customers.csv")
products = pd.read_csv("data/processed/products.csv")

raw = orders.merge(customers, on="customer_id").merge(products, on="product_id")
print(f"Merged: {len(raw)} rows, columns: {list(raw.columns)}")

TEAL = "0F6E62"
wb = load_workbook("excel/regional_performance_summary.xlsx")
ws = wb["Raw Data"]

# clear existing contents completely (old row count may differ from new)
ws.delete_rows(1, ws.max_row)

header_fill = PatternFill(start_color=TEAL, end_color=TEAL, fill_type="solid")
header_font = Font(name="Arial", color="FFFFFF", bold=True, size=10)

for r in dataframe_to_rows(raw, index=False, header=True):
    ws.append(r)
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
for col_cells in ws.columns:
    letter = get_column_letter(col_cells[0].column)
    ws.column_dimensions[letter].width = 14
ws.freeze_panes = "A2"

wb.save("excel/regional_performance_summary.xlsx")
print(f"Done. Raw Data sheet now has {len(raw)} rows. Open the file and check the "
      f"Region Summary / Profit by Region-Year / Sales by CityType-Region tabs recalculated.")
